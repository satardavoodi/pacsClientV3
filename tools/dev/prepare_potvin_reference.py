"""Extract public Potvin supplementary coefficients without executing Excel macros.

Run with openpyxl installed. Input: mmc2.xlsm from PMC5094268 supplements.
The generated model belongs in the ignored local reference runtime directory.
"""
import argparse
import hashlib
import json
from pathlib import Path

WORKBOOK_SHA256 = 'c8a8c9e08d5703923a45eefe2d29ef6dc9c553fc12771f2f35df0af872bfc19f'


def extract(source, destination):
    import openpyxl
    if hashlib.sha256(Path(source).read_bytes()).hexdigest() != WORKBOOK_SHA256:
        raise ValueError('The source workbook is not the reviewed published supplement.')
    workbook = openpyxl.load_workbook(source, data_only=False, keep_vba=False)
    names = {name.lower(): item for name, item in workbook.defined_names.items()}

    def cells(name):
        sheet, address = next(names[name.lower()].destinations)
        value = workbook[sheet][address]
        return [[value.value]] if not isinstance(value, tuple) else [[c.value for c in row] for row in value]

    models = {}
    matrix = workbook['Matrix']
    for row in range(3, 29):
        region = matrix.cell(row, 1).value
        models[region] = {
            'predictors': [str(v).lower() for v in cells('pred_' + region)[0]],
            'coefficients': [r[0] for r in cells('B_' + region)],
            'inverse_design': cells('M_' + region),
            'n': matrix.cell(row, 3).value,
            'mse': matrix.cell(row, 4).value,
            'transform': 'log10' if region.endswith('_log') else 'identity',
        }
    result = {'source': 'https://pmc.ncbi.nlm.nih.gov/articles/PMC5094268/',
              'attribution': 'Potvin, Mouiha, Dieumegarde and Duchesne (2016), CC BY 4.0',
              'workbook_sha256': WORKBOOK_SHA256, 'models': models}
    Path(destination).write_text(json.dumps(result, sort_keys=True, indent=2) + '\n', encoding='utf-8')
    print('Extracted 26 public reference models; no macros executed.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook')
    parser.add_argument('output')
    args = parser.parse_args()
    extract(args.workbook, args.output)
