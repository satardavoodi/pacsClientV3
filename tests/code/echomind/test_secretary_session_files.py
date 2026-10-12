"""Synthetic session source boundaries; never use live User Data."""
import base64
import hashlib
import io
from uuid import uuid4

import pytest


def source(name, raw):
    return dict(name=name, sha256=hashlib.sha256(raw).hexdigest(), data=base64.b64encode(raw).decode())


@pytest.mark.parametrize('mode', ['act','ask','guide','help_ticket'])
def test_sources_reach_server_owned_model_and_receipt(tmp_path, monkeypatch, mode):
    import json
    from modules.ai_imaging.eagle_eye_remote.echomind.hosting import EchoMind
    from modules.ai_imaging.eagle_eye_remote.secretary.service import Secretary, _session_context
    from modules.ai_imaging.eagle_eye_remote.secretary import service
    from PIL import Image
    (tmp_path/'settings.json').write_text('{}')
    echo = EchoMind({'config_dir':str(tmp_path),'history_dir':str(tmp_path/'echo')},['synthetic'])
    host = Secretary({'history_dir':str(tmp_path/'secretary')},echo)
    stream = io.BytesIO()
    Image.new('RGB',(16,16),'green').save(stream,'PNG')
    context = {'id':str(uuid4()),'mode':mode,'title':'Synthetic source review'}
    calls = []
    def completion(**kwargs):
        calls.append(kwargs)
        user = kwargs['messages'][1]['content']
        if len(calls)==1:
            return json.dumps({'modules':['education'],'reason':'Reference request.'})
        assert isinstance(user,list)
        assert any(p['type']=='image_url' for p in user)
        assert 'untrusted reference material' in kwargs['messages'][0]['content']
        if mode=='act':
            return json.dumps({'action':'open_education','entities':{},'confidence':0.9,'needs_confirmation':False,'reason':'Review source course.'})
        result = {'answer':'The supplied reference has been reviewed.'}
        if mode=='help_ticket':
            result['ticket']={'category':'other','description':'Synthetic reference assistance.'}
        return json.dumps(result)
    monkeypatch.setattr(service.llm_client,'gapgpt_chat',completion)
    body = dict(protocol=1,request_id=str(uuid4()),phase='plan',text='Review this supplied reference.',language='en',client_time='2026-10-09T12:00:00+03:30',interaction_mode=mode,conversation=context,session_sources=[source('example.png',stream.getvalue())])
    result = host.process('synthetic',body)
    assert result['session_sources'][0]['status']=='extracted'
    assert result['conversation']['id']==context['id']
    assert _session_context.get() is None
    assert len(calls)==2


def test_limits_and_path_rejection():
    from modules.ai_imaging.eagle_eye_remote.secretary.session_sources import Source, extract
    for name in ('../a.txt','C:\\a.txt','run.exe'):
        with pytest.raises(ValueError):
            Source.model_validate(source(name,b'reference'))
    with pytest.raises(ValueError,match='Duplicate'):
        extract([Source.model_validate(source('a.txt',b'a'))]*2)
    with pytest.raises(ValueError,match='80,000'):
        extract([Source.model_validate(source('a.txt',b'a'*80001))])


def test_mcp_adapter_pending_result_and_stale_session(tmp_path, monkeypatch):
    from types import SimpleNamespace
    from modules.EchoMind.secretary.session_files import SessionFiles
    from modules.EchoMind.secretary import session_files
    from modules.EchoMind.secretary.adapters.session_files_adapter import SessionFilesAdapter
    from modules.EchoMind.secretary.command_envelope import CommandPlan
    session = {'id':str(uuid4()),'mode':'act'}
    monkeypatch.setattr(session_files,'SessionFiles',lambda:SessionFiles(tmp_path))
    memory = SimpleNamespace(get_conversation=lambda:session.copy())
    adapter = SessionFilesAdapter(lambda:memory)
    try:
        receipt = adapter.execute(CommandPlan(action='list_session_files'),{}).data
        adapter._operations[receipt['operation_id']][1].result(timeout=5)
        plan = CommandPlan(action='session_file_status',entities={'operation_id':receipt['operation_id']})
        assert adapter.execute(plan,{}).data['state']=='succeeded'
        session['id']=str(uuid4())
        assert adapter.execute(plan,{}).error_code=='SESSION_CHANGED'
    finally:
        adapter.close()


def test_shared_action_schema_permissions_and_workflow():
    from modules.EchoMind.secretary.command_envelope import validate_action_entities
    from modules.EchoMind.secretary.permissions import classify, READ_ONLY
    from modules.EchoMind.secretary.workflow import make_step
    from modules.EchoMind.secretary.validator import validate_plan as client_validate
    from modules.ai_imaging.eagle_eye_remote.secretary.validation.validator import validate_plan as server_validate
    for action in ('list_session_files','inspect_session_file','session_file_status'):
        entities = {'sha256':'a'*64} if action=='inspect_session_file' else {'operation_id':'a'*32} if action=='session_file_status' else {}
        validate_action_entities(action,entities)
        assert classify(action)==READ_ONLY
        plan = dict(action=action,entities=entities,confidence=0.9,needs_confirmation=False,reason='Inspect supplied references.')
        assert not client_validate(plan)[1]
        assert not server_validate(plan)[1]
    assert make_step('inspect_session_file',{'sha256':'a'*64}).verify.probe_action=='session_file_status'
    with pytest.raises(ValueError):
        validate_action_entities('inspect_session_file',{'sha256':'a'*64,'path':'../private'})


def test_mcp_memory_lookup_does_not_initialize_lazy_database():
    import ast
    from pathlib import Path
    tree = ast.parse(Path('modules/EchoMind/secretary/bus_factory.py').read_text(encoding='utf-8'))
    function = next(n for n in ast.walk(tree) if isinstance(n,ast.FunctionDef) and n.name=='session_memory')
    sentinel = object()
    class Owner:
        def __init__(self):self._memory_store=sentinel
        @property
        def memory_store(self):pytest.fail('GUI dispatch initialized the lazy database')
    scope = {'orchestrator':Owner(),'home_widget':None}
    exec(compile(ast.Module(body=[function],type_ignores=[]),'<memory-lookup>','exec'),scope)
    assert scope['session_memory']() is sentinel


def test_files_button_schedules_folder_work(tmp_path, monkeypatch):
    import sys
    from types import SimpleNamespace
    from PySide6.QtWidgets import QApplication
    from concurrent.futures import ThreadPoolExecutor
    from modules.EchoMind.secretary.session_files import SessionFiles
    from modules.EchoMind.secretary import session_files, remote_planner
    monkeypatch.setitem(sys.modules,'sounddevice',SimpleNamespace())
    from PacsClient.pacs.workstation_ui.home_ui.secretary_button_widget import SecretaryButtonWidget
    app = QApplication.instance() or QApplication([])
    monkeypatch.setattr(remote_planner,'uses_server',lambda:True)
    monkeypatch.setattr(session_files,'SessionFiles',lambda:SessionFiles(tmp_path))
    session = str(uuid4())
    widget = SecretaryButtonWidget()
    widget._ensure_secretary_runtime=lambda:True
    widget._secretary_orchestrator=SimpleNamespace(memory_store=SimpleNamespace(prepare_conversation=lambda mode:{'id':session}))
    queued=[]
    widget._run_worker=lambda work,done,failed:queued.append(work)
    widget._open_session_files()
    assert len(queued)==1
    assert not (tmp_path/session).exists()
    with ThreadPoolExecutor(max_workers=1) as pool:
        assert pool.submit(queued[0]).result()==str(tmp_path/session/'files')
    assert widget.session_files_button.accessibleName()=='Secretary session files'
    widget.resize(290,260)
    widget.layout().activate()
    assert widget.session_files_button.geometry().right() < widget.width()
    assert widget.memory_new_btn.geometry().right() < widget.width()
    widget.close()
    app.processEvents()


def test_client_snapshots_sources_and_requires_matching_receipts(tmp_path, monkeypatch):
    import json
    from concurrent.futures import ThreadPoolExecutor
    from modules.EchoMind.secretary import session_files, remote_planner
    from modules.EchoMind.secretary.session_files import SessionFiles
    store = SessionFiles(tmp_path)
    context = {'id':str(uuid4()),'mode':'ask','title':'Synthetic reference'}
    with ThreadPoolExecutor(max_workers=1) as pool:
        folder = pool.submit(store.folder,context['id']).result()
    (folder/'reference.txt').write_text('Synthetic reference data',encoding='utf-8')
    monkeypatch.setattr(session_files,'SessionFiles',lambda:store)
    bad = [False]
    class Client:
        def json(self,path):
            return {'secretary':{'session_sources':{'version':1}}}
        def open(self,path,body,timeout):
            files = body['session_sources']
            assert base64.b64decode(files[0]['data'])==b'Synthetic reference data'
            result=dict(protocol=1,request_id=body['request_id'],phase=body['phase'],
                route={'modules':['education'],'reason':'Source review.'},plan=None,
                interaction_mode='ask',answer='Reference reviewed.',conversation=context,
                session_sources=[{'name':files[0]['name'],'sha256':'0'*64 if bad[0] else files[0]['sha256'],'status':'extracted'}])
            return io.BytesIO(json.dumps(result).encode())
    monkeypatch.setattr(remote_planner,'Client',Client)
    with ThreadPoolExecutor(max_workers=1) as pool:
        assert pool.submit(remote_planner.request,'plan','Summarize my reference.',interaction_mode='ask',conversation=context).result()['answer']=='Reference reviewed.'
        bad[0]=True
        with pytest.raises(remote_planner.RemotePlanningError,match='acknowledge'):
            pool.submit(remote_planner.request,'plan','Summarize my reference.',interaction_mode='ask',conversation=context).result()


def test_session_folder_and_new_isolation(tmp_path):
    from concurrent.futures import ThreadPoolExecutor
    from modules.EchoMind.secretary.session_files import SessionFiles
    store = SessionFiles(tmp_path)
    first, second = str(uuid4()), str(uuid4())
    with ThreadPoolExecutor(max_workers=1) as pool:
        (pool.submit(store.folder,first).result() / 'article.txt').write_text('Synthetic reference')
        assert len(pool.submit(store.snapshot,first).result()) == 1
        assert pool.submit(store.snapshot,second).result() == []
        with pytest.raises(ValueError):
            pool.submit(store.folder,'../outside').result()


def test_server_reads_actual_pixels_and_rejects_digest():
    from PIL import Image
    from modules.ai_imaging.eagle_eye_remote.secretary.session_sources import Source, extract
    stream = io.BytesIO()
    Image.new('RGB', (32, 24), 'blue').save(stream, 'JPEG')
    raw = stream.getvalue()
    item = dict(name='reference.jpg', sha256=hashlib.sha256(raw).hexdigest(), data=base64.b64encode(raw).decode())
    parts, receipt = extract([Source.model_validate(item)])
    assert parts[1]['type'] == 'image_url'
    assert base64.b64decode(parts[1]['image_url']['url'].split(',')[1])
    assert receipt[0]['sha256'] == item['sha256']
    with pytest.raises(ValueError, match='digest'):
        extract([Source.model_validate(dict(item, sha256='0'*64))])


def test_pptx_text_and_pdf_pages():
    from pptx import Presentation
    from pypdf import PdfWriter
    from modules.ai_imaging.eagle_eye_remote.secretary.session_sources import Source, extract
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = 'Synthetic course source'
    ppt = io.BytesIO()
    deck.save(ppt)
    pdf = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.write(pdf)
    def source(name, raw):
        return Source(name=name, sha256=hashlib.sha256(raw).hexdigest(), data=base64.b64encode(raw).decode())
    parts, receipt = extract([source('slides.pptx', ppt.getvalue()), source('scan.pdf', pdf.getvalue())])
    assert any('Synthetic course source' in p.get('text','') for p in parts)
    assert any(p['type'] == 'image_url' for p in parts)
    assert [r['units'] for r in receipt] == [1, 1]
